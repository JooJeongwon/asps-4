"""Render a cited deck once to PDF, then use those pages in PowerPoint.

PPTX pages are images (speaker notes remain text). This makes Korean typography
and PDF/PPTX layout identical without a separate Office installation.
"""
from __future__ import annotations

import html
import json
from pathlib import Path
import re

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches


def validate(data, sources, seconds):
    slides = data.get("slides")
    if not isinstance(slides, list) or not 3 <= len(slides) <= 24:
        raise ValueError("expected 3–24 slides")
    source_map = {s["id"]: s for s in sources}
    if sum(s["seconds"] for s in slides) != seconds:
        raise ValueError("slide timing does not equal presentation target")
    for slide in slides:
        if not 1 <= slide["seconds"] <= 180 or not slide["notes"].strip():
            raise ValueError("missing speaker notes or invalid timing")
        if len(slide["title"]) > 32 or len(slide["message"]) > 70 or not 1 <= len(slide["bullets"]) <= 4:
            raise ValueError("slide text is too dense")
        for bullet in slide["bullets"]:
            if bullet["kind"] not in {"fact", "self_report", "proposal", "unknown"} or len(bullet["text"]) > 100:
                raise ValueError("invalid bullet")
            if bullet["kind"] in {"fact", "self_report"} and not bullet["evidence"]:
                raise ValueError("factual bullet lacks evidence")
            for e in bullet["evidence"]:
                if e["source"] not in source_map or not e["location"].strip():
                    raise ValueError("unknown source or missing location")
                if e["source"] == "S1":
                    page_labels = re.findall(r"\bp{1,2}\.\s*(\d+(?:\s*[-–—,]\s*\d+)*)", e["location"], re.I)
                    nums = re.findall(r"\d+", " ".join(page_labels))
                    if not nums or any(not 1 <= int(n) <= 22 for n in nums):
                        raise ValueError("invalid report page")
    speech_chars = sum(len(re.sub(r"\s", "", s["notes"])) for s in slides)
    if speech_chars > seconds * 7:
        raise ValueError("script exceeds estimated speech rate of 7 characters/sec")
    return speech_chars


def render(data, sources, run, seconds, basename="neofirm-final", status="자동 생성 / 자료·가정 구분"):
    run = Path(run)
    run.mkdir(parents=True, exist_ok=True)
    speech_chars = validate(data, sources, seconds)
    slides = data["slides"]
    width, height = 960, 540
    labels = {"fact": "근거", "self_report": "자기보고", "proposal": "설계 가정", "unknown": "미확인"}
    def esc(text):
        # AppleGothic lacks U+30FB; use the equivalent supported middle dot.
        return html.escape(text.replace("・", "·"))
    font = Path("/System/Library/Fonts/Supplemental/AppleGothic.ttf")
    css = """* {box-sizing:border-box} body {font-family: sans-serif; color:#ecf1fa; margin:0}
    .top {font-size:12px; color:#77dacf; letter-spacing:2px; margin-bottom:24px}
    h1 {font-size:34px; line-height:1.25; margin:0 0 16px}
    .message {font-size:23px; line-height:1.4; color:#a9bdd7; margin-bottom:26px}
    .card {background:#152439; border-left:3px solid #77dacf; padding:11px 16px; margin:0 0 11px}
    .label {font-size:11px; color:#77dacf; margin-bottom:5px}
    .text {font-size:19px; line-height:1.4}
    """
    archive = None
    if font.exists():
        archive = pymupdf.Archive(str(font.parent))
        css += "@font-face {font-family: Korean; src:url(AppleGothic.ttf)} body {font-family:Korean}"
    pdf = pymupdf.open()
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333333), Inches(7.5)
    thumbs, quality, outline, script = [], [], [], []
    source_map = {s["id"]: s for s in sources}
    for i, slide in enumerate(slides, 1):
        cards, citations = [], []
        for bullet in slide["bullets"]:
            refs = "; ".join(f"{e['source']} {e['location']}" for e in bullet["evidence"])
            label = labels[bullet["kind"]] + (" · " + refs if refs else "")
            cards.append(f'<div class="card"><div class="label">{esc(label)}</div><div class="text">{esc(bullet["text"])}</div></div>')
            citations.extend(bullet["evidence"])
        body = f'<div class="top">N E O F I R M / {i:02d}</div><h1>{esc(slide["title"])}</h1><div class="message">{esc(slide["message"])}</div>' + "".join(cards)
        page = pdf.new_page(width=width, height=height)
        page.draw_rect(page.rect, color=None, fill=(0.035, 0.075, 0.13))
        spare, scale = page.insert_htmlbox(pymupdf.Rect(42, 25, 918, 498), body, css=css, archive=archive, scale_low=1)
        if spare < 0 or scale < 0.99:
            raise ValueError(f"slide {i}: clipped or undersized text")
        footer = f'{esc(status)} · {slide["seconds"]}초 · {i}/{len(slides)}'
        page.insert_htmlbox(pymupdf.Rect(42, 506, 918, 535), footer,
                            css=css + "body {font-size:10px;color:#8ca1bd}", archive=archive, scale_low=1)
        png = run / "slides" / f"{i:02d}.png"
        png.parent.mkdir(parents=True, exist_ok=True)
        page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(png)
        target = presentation.slides.add_slide(presentation.slide_layouts[6])
        target.shapes.add_picture(str(png), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
        bibliography = "\n".join(f"{e['source']} {e['location']} — " + source_map[e["source"]].get("url", source_map[e["source"]].get("path", "")) for e in citations)
        target.notes_slide.notes_text_frame.text = slide["notes"] + "\n\n" + bibliography
        extracted = page.get_text()
        if "\ufffd" in extracted or any(char not in extracted for char in slide["title"] if "가" <= char <= "힣"):
            raise ValueError(f"slide {i}: missing Korean glyphs")
        quality.append({"slide": i, "scale": scale, "spare_height": spare, "korean_text_check": "passed"})
        outline.append(f"## {i}. {slide['title']} ({slide['seconds']}초)\n\n{slide['message']}\n\n"
                       + "\n".join(f"- [{labels[b['kind']]}] {b['text']} " + "; ".join(f"[{e['source']} {e['location']}]" for e in b["evidence"]) for b in slide["bullets"]))
        script.append(f"## {i}. {slide['title']} ({slide['seconds']}초)\n\n{slide['notes']}\n\n출처:\n{bibliography}")
        with Image.open(png) as original:
            thumb = original.copy()
        thumb.thumbnail((480, 270))
        thumbs.append(thumb)
    (run / "slides-outline.md").write_text("# 슬라이드 구성\n\n" + "\n\n".join(outline), encoding="utf-8")
    pdf.save(run / f"{basename}.pdf", garbage=4, deflate=True)
    pdf.close()
    presentation.save(run / f"{basename}.pptx")
    reloaded = Presentation(run / f"{basename}.pptx")
    if len(reloaded.slides) != len(slides):
        raise ValueError("PowerPoint reopen check failed")
    sheet = Image.new("RGB", (960, 290 * ((len(thumbs) + 1) // 2)), "#ffffff")
    draw = ImageDraw.Draw(sheet)
    for i, thumb in enumerate(thumbs):
        x, y = (i % 2) * 480, (i // 2) * 290
        sheet.paste(thumb, (x, y))
        draw.text((x + 10, y + 272), f"{i+1:02d}", fill="#172333")
    sheet.save(run / "contact-sheet.png")
    (run / "speaker-notes.md").write_text("# 발표 대본\n\n" + "\n\n".join(script), encoding="utf-8")
    (run / "quality.json").write_text(json.dumps({"slides": quality, "seconds": seconds,
        "speech_characters": speech_chars, "speech_rate_assumption": "7 Korean characters/sec; rehearsal needed",
        "pptx_reopen": "passed", "pptx_layout": "PDF raster pages; text is not editable",
        "citation_check": "source IDs and locations validated; semantic citation review still needed",
        "status": status}, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("run", type=Path)
    parser.add_argument("--minutes", type=int, default=7)
    parser.add_argument("--draft", action="store_true")
    args = parser.parse_args()
    render(json.loads((args.run / "slides.json").read_text()), json.loads((args.run / "sources.json").read_text()),
           args.run, args.minutes * 60, basename="neofirm-review" if args.draft else "neofirm-final",
           status="검토용 초안 / NotebookLM 인증 대기" if args.draft else "자동 생성 / 자료·가정 구분")
